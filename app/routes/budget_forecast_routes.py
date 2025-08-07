from flask import Blueprint, request, jsonify, send_file
import pandas as pd
from datetime import datetime
import os
from flask_cors import CORS
from langchain.chains import LLMChain
from langchain_community.llms import Ollama
from langchain_core.prompts import PromptTemplate

budget_forecast_routes = Blueprint('budget_forecast_routes', __name__)
CORS(budget_forecast_routes)

@budget_forecast_routes.route("/predict-only", methods=["POST"])
def predict_only():
    try:
        file = request.files["file"]
        df = pd.read_excel(file)

        predictions_df = generate_budget_predictions(df)
        output_path = save_predictions_to_excel(predictions_df)

        filename = os.path.basename(output_path)
        return jsonify({"excel_url": f"/api/reports/{filename}", "stats": extract_stats(predictions_df)})

    except Exception as e:
        print(e)
        return jsonify({"error": str(e)}), 500


@budget_forecast_routes.route("/generate-report", methods=["POST"])
def generate_report():
    try:
        # Load file again
        data = request.get_json()
        filename = data["filename"]

        file_path = os.path.join(os.path.dirname(__file__), 'reports', filename)
        df = pd.read_excel(file_path)

        # ⬇️ Compute meaningful stats
        summary = df.groupby('Année')[['Investissement', 'Fonctionnement', 'Services', 'Formation', 'Consommation', 'Écart']].sum().reset_index()

        # Compute year-over-year % changes
        pct_changes = summary.set_index('Année').pct_change().fillna(0) * 100
        pct_changes = pct_changes.round(2)

        stats_summary = ""
        for year in summary['Année'].tolist():
            row = summary[summary['Année'] == year].iloc[0]
            stats_summary += f"\n📊 Année {year}:\n"
            for col in ['Investissement', 'Fonctionnement', 'Services', 'Formation', 'Consommation', 'Écart']:
                stats_summary += f"  - {col}: {row[col]:,.0f} TND\n"

        stats_summary += "\n🔁 Pourcentages d'évolution:\n"
        for year in pct_changes.index.tolist()[1:]:
            stats_summary += f"\n📉 Évolution de {year} vs {year - 1}:\n"
            for col in pct_changes.columns:
                stats_summary += f"  - {col}: {pct_changes.loc[year][col]:.2f}%\n"

        # 🧠 Use real stats
        prompt = PromptTemplate(
            input_variables=["stats"],
            template=(
                "Tu es un expert en analyse financière. Génère un rapport clair, structuré et professionnel en **français** à partir des données suivantes :\n\n"
                "{stats}\n\n"
                "Le rapport doit contenir :\n"
                "- Un **titre**.\n"
                "- Une **introduction** expliquant le contexte global.\n"
                "- Une **analyse année par année** avec les montants et pourcentages d'évolution.\n"
                "- Une section **alertes et recommandations stratégiques**.\n"
                "- Des **sauts de ligne** clairs entre les parties pour une lecture fluide.\n"
                "- Un **langage formel et analytique**, à destination de la direction financière.\n\n"
                "Structure le rapport avec des paragraphes bien séparés, sans créer de liste à puces sauf si nécessaire."
            )
        )
        llm = Ollama(model="mistral")
        chain = LLMChain(llm=llm, prompt=prompt)
        report = chain.run(stats=stats_summary)

        return jsonify({"report": report})

    except Exception as e:
        print(e)
        return jsonify({"error": str(e)}), 500

@budget_forecast_routes.route('/api/reports/<filename>', methods=['GET'])
def download_report(filename):
    try:
        return send_file(os.path.join(os.path.dirname(__file__), 'reports', filename), as_attachment=True)
    except Exception as e:
        return jsonify({"error": str(e)}), 404


def generate_budget_predictions(df):
    current_year = datetime.now().year
    groupements_structure = df[['Direction', 'Rubrique', 'Groupement']].drop_duplicates()

    aggregated = df.groupby(['Année', 'Direction', 'Rubrique', 'Groupement'])[ 
        ['Investissement', 'Fonctionnement', 'Services', 'Formation', 'Consommation']
    ].sum().reset_index()

    historical_deltas = []
    for (direction, rubrique, groupement), group in aggregated.groupby(['Direction', 'Rubrique', 'Groupement']):
        group_sorted = group.sort_values('Année')
        deltas = group_sorted[['Investissement', 'Fonctionnement', 'Services', 'Formation']].diff().dropna()
        if not deltas.empty:
            avg_delta = deltas.mean()
            historical_deltas.append({
                'Direction': direction,
                'Rubrique': rubrique,
                'Groupement': groupement,
                'Δ_Investissement': avg_delta['Investissement'],
                'Δ_Fonctionnement': avg_delta['Fonctionnement'],
                'Δ_Services': avg_delta['Services'],
                'Δ_Formation': avg_delta['Formation']
            })

    delta_df = pd.DataFrame(historical_deltas)

    last_known = aggregated[aggregated['Année'] == current_year - 1]
    current_predictions = last_known.merge(delta_df, on=['Direction', 'Rubrique', 'Groupement'], how='left')

    for col in ['Investissement', 'Fonctionnement', 'Services', 'Formation']:
        current_predictions[col] = current_predictions[col] + current_predictions[f'Δ_{col}']

    current_predictions['Année'] = current_year
    current_predictions['Consommation'] = current_predictions[['Investissement', 'Fonctionnement', 'Services', 'Formation']].sum(axis=1) * 0.8
    current_predictions['Écart'] = current_predictions[['Investissement', 'Fonctionnement', 'Services', 'Formation']].sum(axis=1) - current_predictions['Consommation']

    forecast = []
    latest = current_predictions.copy()
    for i in range(1, 3):
        next_year = current_year + i
        next_pred = latest.copy()
        next_pred['Année'] = next_year
        for col in ['Investissement', 'Fonctionnement', 'Services', 'Formation']:
            next_pred[col] = next_pred[col] + next_pred[f'Δ_{col}']
        next_pred['Consommation'] = next_pred[['Investissement', 'Fonctionnement', 'Services', 'Formation']].sum(axis=1) * 0.8
        next_pred['Écart'] = next_pred[['Investissement', 'Fonctionnement', 'Services', 'Formation']].sum(axis=1) - next_pred['Consommation']
        forecast.append(next_pred)
        latest = next_pred.copy()

    forecast_df = pd.concat([current_predictions] + forecast, ignore_index=True)
    return forecast_df


def save_predictions_to_excel(df):
    reports_dir = os.path.join(os.path.dirname(__file__), 'reports')
    os.makedirs(reports_dir, exist_ok=True)

    filename = f"prediction_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    filepath = os.path.join(reports_dir, filename)

    df.to_excel(filepath, index=False)
    return filepath


def extract_stats(df):
    try:
        summary = df.groupby("Année")[['Investissement', 'Fonctionnement', 'Services', 'Formation', 'Consommation', 'Écart']].sum().round(1)
        return summary.to_string()
    except Exception as e:
        return "No summary available."
def remove_emojis(text):
    # Remove all non-Latin-1 characters (emojis, special unicode)
    return text.encode('latin-1', 'ignore').decode('latin-1')


from sqlalchemy import text
from app.extensions import db

@budget_forecast_routes.route("/all-years-report", methods=["GET"])
def get_consolidated_report_all_years():
    sql = text("""
        SELECT 
            b.year,
            d.name AS direction,
            r.name AS rubrique,
            g.name AS groupement,
            g.budget_alloue,
            SUM(CASE WHEN pd.type = 'Investissement' THEN pd.montant ELSE 0 END) AS investissement,
            SUM(CASE WHEN pd.type = 'Fonctionnement' THEN pd.montant ELSE 0 END) AS fonctionnement,
            SUM(CASE WHEN pd.type = 'Services' THEN pd.montant ELSE 0 END) AS services,
            SUM(CASE WHEN pd.type = 'Formation' THEN pd.montant ELSE 0 END) AS formation,
            g.budget_consomme,
            g.ecart
        FROM Budgets b
        JOIN Groupements g ON g.budget_id = b.id
        JOIN Rubriques r ON r.id = g.rubrique_id
        JOIN Directions d ON d.id = r.direction_id
        LEFT JOIN ProjetDetails pd ON pd.groupement_id = g.id
        GROUP BY b.year, d.name, r.name, g.name, g.budget_alloue, g.budget_consomme, g.ecart
        ORDER BY b.year, d.name, r.name, g.name
    """)

    results = db.session.execute(sql).fetchall()
    return jsonify([dict(row._mapping) for row in results])

@budget_forecast_routes.route("/predicted-report-data", methods=["POST"])
def get_predicted_report_data():
    try:
        file = request.files["file"]
        df = pd.read_excel(file)

        predictions_df = generate_budget_predictions(df)

        # 🔁 Normalize column names for Angular compatibility
        renamed = predictions_df.rename(columns={
            'Année': 'year',
            'Direction': 'direction',
            'Rubrique': 'rubrique',
            'Groupement': 'groupement',
            'Investissement': 'investissement',
            'Fonctionnement': 'fonctionnement',
            'Services': 'services',
            'Formation': 'formation',
            'Consommation': 'budget_consomme',
            'Écart': 'ecart'
        })

        return jsonify(renamed.to_dict(orient="records"))

    except Exception as e:
        print(e)
        return jsonify({"error": str(e)}), 500


@budget_forecast_routes.route("/generate-full-pdf", methods=["POST"])
def generate_full_pdf():
    try:
        data = request.get_json()
        filename = data["filename"]
        report = data["report"]

        # Load Excel data
        file_path = os.path.join(os.path.dirname(__file__), 'reports', filename)
        df = pd.read_excel(file_path)

        # Generate PDF
        from fpdf import FPDF

        class PDF(FPDF):
            def header(self):
                self.set_font('Arial', 'B', 14)
                self.cell(0, 10, 'Rapport Budgétaire Complet', 0, 1, 'C')
                self.ln(10)

        pdf = PDF()
        pdf.add_page()
        pdf.set_font("Arial", size=10)

        # Table part
        pdf.cell(0, 10, "Résumé des Budgets", ln=True)
        pdf.ln(2)
        for _, row in df.iterrows():
            line = f"{row['Année']} | {row['Direction']} | {row['Rubrique']} | {row['Groupement']} | {row['Investissement']:.0f} | {row['Fonctionnement']:.0f}"
            pdf.cell(0, 8, line, ln=True)

        # AI Report part
        pdf.add_page()
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(0, 10, "Analyse Générative IA", ln=True)
        pdf.ln(5)
        pdf.set_font("Arial", '', 10)

        for line in report.split("\n"):
            pdf.multi_cell(0, 8, line)

        # Save
        output_path = os.path.join(os.path.dirname(__file__), 'reports', f"rapport_complet_{datetime.now().strftime('%Y%m%d%H%M%S')}.pdf")
        pdf.output(output_path)

        return jsonify({"pdf_url": f"/api/reports/{os.path.basename(output_path)}"})

    except Exception as e:
        print(e)
        return jsonify({"error": str(e)}), 500

@budget_forecast_routes.route('/export-excel-report', methods=['POST'])
def export_excel_report():
    from openpyxl import Workbook
    from openpyxl.utils.dataframe import dataframe_to_rows

    try:
        data = request.get_json()
        filename = data["filename"]
        report_text = data["report"]

        file_path = os.path.join(os.path.dirname(__file__), 'reports', filename)
        df = pd.read_excel(file_path)

        wb = Workbook()
        ws_summary = wb.active
        ws_summary.title = "Budget Par Année"

        for year, group in df.groupby("Année"):
            ws_summary.append([f"📅 Année {year}"])
            columns = ['Direction', 'Rubrique', 'Groupement', 'Investissement',
                       'Fonctionnement', 'Services', 'Formation', 'Budget Alloué',
                       'Consommation', 'Écart']
            ws_summary.append(columns)
            for _, row in group.iterrows():
                ws_summary.append([
                    row.get('Direction', ''),
                    row.get('Rubrique', ''),
                    row.get('Groupement', ''),
                    row.get('Investissement', 0),
                    row.get('Fonctionnement', 0),
                    row.get('Services', 0),
                    row.get('Formation', 0),
                    row.get('Investissement', 0) + row.get('Fonctionnement', 0) +
                    row.get('Services', 0) + row.get('Formation', 0),
                    row.get('Consommation', 0),
                    row.get('Écart', 0),
                ])
            ws_summary.append([])

        # Add AI report in another sheet
        ws_report = wb.create_sheet(title="Rapport IA")
        for i, line in enumerate(report_text.splitlines(), start=1):
            ws_report.cell(row=i, column=1, value=line)

        output_filename = f"Export_Budget_Report_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        output_path = os.path.join(os.path.dirname(__file__), 'reports', output_filename)
        wb.save(output_path)

        return jsonify({"excel_export_url": f"/api/reports/{output_filename}"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    